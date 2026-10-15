"""Build the Model B brief and methods report using the bundled document runtime.

This only creates documents and figures. It never starts a GPU benchmark.
"""
from datetime import datetime, timezone
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/briefs/2026-10-15-model-b"
WORK = ROOT / ".benchmark-results/documents/2026-10-15-model-b"
BASE_URL = "https://github.com/anik1617/Ratcliff/blob/2b75421174c34bcde708e8d1cecd47df3934f151/"
FONT = Path("C:/Windows/Fonts")
BLACK = RGBColor(0, 0, 0)


def font(size, bold=False):
    return ImageFont.truetype(str(FONT / ("calibrib.ttf" if bold else "calibri.ttf")), size)


def centered(draw, box, text, size=31, bold=False, fill="#111111"):
    x0, y0, x1, y1 = box
    draw.multiline_text(((x0+x1)/2, (y0+y1)/2), text, font=font(size, bold),
                        fill=fill, anchor="mm", align="center", spacing=7)


def make_diagrams():
    image = Image.new("RGB", (1800, 240), "white")
    draw = ImageDraw.Draw(image)
    labels = ["Simulate 32 steps", "Record the first decision", "Continue unfinished trials"]
    for i, text in enumerate(labels):
        x = 8 + i * 604
        draw.rounded_rectangle((x, 20, x+550, 160), radius=13, fill="#edf3f2", outline="#41776b", width=3)
        centered(draw, (x, 20, x+550, 160), text, size=33, bold=True)
        if i < 2:
            draw.line((x+556, 90, x+588, 90), fill="#303a40", width=5)
            draw.polygon(((x+591, 90), (x+576, 79), (x+576, 101)), fill="#303a40")
    centered(draw, (0, 175, 1800, 232), "Every intermediate time step is checked inside each block", size=30)
    image.save(WORK / "early-stop-flow.png")

    image = Image.new("RGB", (1800, 650), "white")
    draw = ImageDraw.Draw(image)
    centered(draw, (0, 0, 1800, 65), "Stop generating fields after a trial has already decided", 40, True)
    steps = 13
    start, width, gap = 380, 88, 12
    centered(draw, (0, 115, 340, 220), "Original\nfull horizon", 32, True)
    centered(draw, (0, 310, 340, 415), "Improvement 1\nearly stopping", 32, True)
    for i in range(steps):
        x = start + i*(width+gap)
        draw.rectangle((x, 110, x+width, 220), fill="#778493", outline="white", width=2)
        if i < 3:
            draw.rectangle((x, 305, x+width, 415), fill="#377f70", outline="white", width=2)
        else:
            draw.rectangle((x, 305, x+width, 415), fill="#f1f3f4", outline="#c8ced3", width=2)
    x = start + 3*(width+gap) - gap/2
    draw.line((x, 83, x, 450), fill="#a74646", width=4)
    centered(draw, (start, 452, 900, 530), "Decision already recorded", 30, True, "#9a3838")
    centered(draw, (900, 452, 1770, 530), "Unused future fields are avoided", 30, True)
    centered(draw, (40, 550, 1760, 640), "Illustration for one trial  |  up to 400 steps  |  unfinished trials are repacked into smaller batches", 28)
    image.save(WORK / "work-avoided.png")


def setup(compact=False):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(.62 if compact else .72)
    sec.left_margin = sec.right_margin = Inches(.68 if compact else .75)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Caption"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = BLACK
        style.font.underline = False
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.05 if compact else 1.1
    normal.paragraph_format.space_after = Pt(4 if compact else 7)
    doc.styles["Title"].font.size = Pt(20 if compact else 24)
    doc.styles["Title"].font.bold = True
    doc.styles["Title"].paragraph_format.space_after = Pt(5)
    doc.styles["Heading 1"].font.size = Pt(11.5 if compact else 15)
    doc.styles["Heading 1"].font.bold = True
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(7 if compact else 10)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(4 if compact else 6)
    doc.styles["Caption"].font.size = Pt(9)
    doc.styles["Caption"].font.italic = False
    doc.core_properties.author = "Aniket Chaudhari"
    doc.core_properties.subject = "Model B simulation acceleration and validation"
    doc.core_properties.keywords = "Model B, spatial diffusion, early stopping, JAX, validation"
    doc.core_properties.created = datetime(2026, 10, 15, tzinfo=timezone.utc)
    # The bundled base template contains a blue rule in its Title style.
    for border in list(doc.styles.element.xpath(".//w:pBdr")):
        border.getparent().remove(border)
    if not compact:
        footer = sec.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        label = footer.add_run("Model B research report  |  ")
        label.font.size = Pt(9)
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)
    return doc


def p(doc, text, bold=False, size=None):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    return para


def heading(doc, text):
    doc.add_heading(text, level=1)


def new_page(doc, text):
    para = doc.add_heading(text, level=1)
    para.paragraph_format.page_break_before = True


def picture(doc, path, width=7):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(4)
    para.add_run().add_picture(str(path), width=Inches(width))
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER


def table(doc, headers, rows, widths, compact=False):
    tab = doc.add_table(rows=1, cols=len(headers))
    tab.autofit = False
    for col, width in zip(tab.columns, widths):
        col.width = Inches(width)
    for i, text in enumerate(headers):
        tab.rows[0].cells[i].text = str(text)
    for row in rows:
        cells = tab.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = str(text)
    for ri, row in enumerate(tab.rows):
        trpr = row._tr.get_or_add_trPr()
        no_split = OxmlElement("w:cantSplit")
        trpr.append(no_split)
        if ri == 0:
            trpr.append(OxmlElement("w:tblHeader"))
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                el = OxmlElement("w:" + edge)
                for key, value in (("val", "single"), ("sz", "5"), ("color", "D9D9D9")):
                    el.set(qn("w:"+key), value)
                borders.append(el)
            tcpr.append(borders)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "36434B" if ri == 0 else "F3F5F6" if ri % 2 == 0 else "FFFFFF")
            tcpr.append(shading)
            margins = OxmlElement("w:tcMar")
            for side, val in (("top", "70"), ("bottom", "70"), ("left", "90"), ("right", "90")):
                el = OxmlElement("w:"+side)
                el.set(qn("w:w"), val)
                el.set(qn("w:type"), "dxa")
                margins.append(el)
            tcpr.append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.02
                if ci > 0:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    run.font.size = Pt(9.3 if compact else 10)
                    run.font.bold = ri == 0
                    run.font.color.rgb = RGBColor(255,255,255) if ri == 0 else BLACK
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return tab


def link(doc, label, path, base_url=BASE_URL):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(5)
    rel = para.part.relate_to(base_url + path,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "24596B")
    props.append(color)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    para._p.append(hyperlink)
    return para


def brief():
    doc = setup(True)
    doc.add_paragraph("Model B simulation speedup", style="Title")
    p(doc, "Research brief for Dr Ratcliff  |  Aniket Chaudhari  |  15 October 2026", size=9)
    p(doc, "We added blockwise early stopping to Gowrish's existing JAX and K-L implementation. "
      "Model B forward simulation was 10.9 times faster on the tested H100 workload. "
      "The simulation result is reproducible; faster fitting at acceptable accuracy is still unproven.")
    heading(doc, "What we changed")
    p(doc, "The existing code generated noise fields through the full 400-step horizon. "
      "We generate 32-step blocks, check every step for the first decision, save finished trials, "
      "and repack only unfinished trials for the next block. The K-L approximation is unchanged.")
    picture(doc, WORK / "early-stop-flow.png", 7.05)
    heading(doc, "Measured results")
    table(doc, ["Comparison", "Before", "After", "Speedup"], [
        ["June JAX objective  FFT to K-L", "10.66 s", "4.71 s", "2.26x"],
        ["October H100 simulation  K-L to early stop", "3.9447 s", "0.3618 s", "10.90x"],
        ["Laptop RTX 3060 simulation  K-L to early stop", "90.85 s", "8.60 s", "10.56x"],
    ], [3.6, 1.1, 1.1, 1.25], True)
    p(doc, "These are different workloads and hardware comparisons; the speedup factors must not be multiplied. "
      "Historical Fortran timings provide context, but no exact overall Fortran-to-new-fit multiplier is established.", size=9)
    heading(doc, "How we confirmed the simulation result")
    p(doc, "H100: 9,000 trials per condition, two conditions, seven warmed timed calls with GPU completion waits. "
      "Across 12 random seeds, all 216,000 paired trial comparisons matched the full-horizon blocked control "
      "exactly in both choice and reaction time. Independent original-versus-new distribution screens passed "
      "at this parameter setting. The 95% timing bootstrap interval was 10.83-10.91x.")
    p(doc, "Laptop: three parameter settings and three seeds at 512 trials per condition gave 9,216 further "
      "exact paired comparisons and 8.94-10.78x simulation speedups. Settings, seeds, source fingerprints "
      "and hashed raw outputs are preserved.")
    heading(doc, "What remains before publication")
    p(doc, "All eight laptop fits missed the accuracy target. CPU replay confirmed flat search plateaus, and "
      "diagnostics found noisy 512-trial scores and K-L coefficient reassignment under tiny parameter changes. "
      "June's recovery example started at 3.32% error, already below today's 5% recovery threshold. "
      "Repairing these fitting limitations still requires recovery tests.")
    p(doc, "We can start a Model B simulation-methods draft now. Before finalizing its claims, agree scientific "
      "tolerances, expand parameter and boundary-case checks, and repeat matched benchmarks from a frozen version. "
      "A claim about faster accurate fitting additionally needs successful, quality-matched recovery studies.")
    p(doc, "Evidence and methods are documented in the accompanying report. The current result supports tested "
      "simulation acceleration, not universal equivalence or a new diffusion model.", size=9)
    doc.save(OUT / "Model_B_One_Page_Brief.docx")


def report():
    doc = setup()
    doc.add_paragraph("Model B methods and validation", style="Title")
    p(doc, "Prepared for Dr Ratcliff by Aniket Chaudhari  |  15 October 2026", size=9.5)
    heading(doc, "Purpose and main result")
    p(doc, "This report explains the existing Model B acceleration, our early-stopping extension, the measured "
      "gains and the evidence still needed for a paper. The strongest new result is approximately 10.9-fold "
      "faster warmed forward simulation on an H100 at the tested parameters. That result is supported by "
      "a paired correctness control and a second GPU study. Accurate fitting faster has not yet been established.")
    heading(doc, "The original work")
    p(doc, "Model B accumulates evidence on a 100 by 160 spatial field over up to 400 time steps and returns "
      "one of five response categories plus reaction time. Gowrish's implementation moved the calculation into "
      "JAX, batched GPU work, and introduced a truncated Karhunen-Loeve or K-L noise generator. "
      "It generates correlated fields using a smaller spectral basis and matrix multiplication. "
      "Caching and fixed-width padding keep repeated fitting evaluations from rebuilding and recompiling unnecessarily. [1, 7]")
    table(doc, ["June recorded comparison", "FFT", "Cached K-L", "Ratio"], [
        ["Two-condition objective evaluation", "10.66 s", "4.71 s", "2.26x"],
        ["Whole synthetic recovery run", "30.3 min", "20.1 min", "1.51x"],
        ["Mean active-parameter error", "3.8%", "2.98%", "Different estimates"],
    ], [3.15, 1.15, 1.25, 1.45])
    p(doc, "These historical runs used 9,000 fitting trials per condition, synthetic observations generated "
      "with 1,024 trials per condition and a 100-iteration budget. The K-L run made 256 objective evaluations "
      "versus 171 for FFT, and stopped at its iteration limit. Ten of twelve active parameters were within "
      "7% of truth; the worst error was 9.5%. This was useful recovery in one example, not a convergence guarantee. [1]")
    heading(doc, "What the Fortran comparisons mean")
    p(doc, "Archived Fortran timings were about 11 seconds per simulator call on six MPI nodes and 36 seconds "
      "on one node. A two-condition Fortran objective of about 22 seconds is an estimate obtained by doubling "
      "the six-node call time. Relative to the 4.71-second JAX objective, this suggests roughly 4.7-fold "
      "acceleration across different hardware. It is not a newly measured, matched end-to-end comparison. "
      "No exact whole-fit ratio follows from the description of old fits as taking hours. [1]")
    p(doc, "Our October experiment measures an additional change to the existing K-L simulator. "
      "Its speedup cannot be multiplied by the historical objective or Fortran ratios.", bold=True)

    new_page(doc, "How improvement 1 works")
    p(doc, "A trial is finished once the first threshold crossing determines its choice and reaction time. "
      "Calculating future fields for that trial cannot change that already-recorded result. "
      "The old vectorized path nevertheless generated the complete horizon for each batch. "
      "Our extension avoids much of that unused work while retaining GPU batching. [2, 6]")
    picture(doc, WORK / "work-avoided.png", 7)
    p(doc, "The diagram is illustrative. A block contains 32 time steps; it is not a single threshold check.", size=9)
    for title, text in [
        ("Generate a short block", "Create the next block of correlated noise and evidence increments for unfinished trials using the existing K-L representation."),
        ("Check every intermediate step", "Accumulate evidence and find the earliest crossing within the block. Checking only the block endpoint could miss a decision and would be scientifically incorrect."),
        ("Save completed trials", "Store their first choice and reaction time. They do not need future blocks. Trials without a crossing continue, up to the original horizon and fallback rules."),
        ("Repack the remaining work", "Compact unfinished trials into smaller supported batches so the GPU does less padded work. Random keys depend on absolute trial and block identifiers, rather than a trial's temporary batch position."),
    ]:
        p(doc, title, bold=True)
        p(doc, text)
    heading(doc, "What stays scientifically separate")
    p(doc, "Early stopping does not tighten the K-L truncation. At the reference correlation parameter, "
      "1,325 complex modes capture approximately 99% of variance; the usual padded capacity is 2,000 complex "
      "modes. This is the existing approximation. Exact agreement with a blocked control tests the stopping "
      "implementation, not whether truncated K-L is identical to the full FFT covariance. [1, 7]")
    p(doc, "The original and blocked implementations use different random-number layouts. Equal root seeds "
      "therefore do not require identical original-versus-new sample paths. We use exact paired checks "
      "against a full-horizon blocked control and separate statistical checks against the original implementation.")

    new_page(doc, "Measured simulation performance")
    table(doc, ["Workload and device", "Original K-L", "Early stop", "Speedup"], [
        ["H100 SXM 80 GB  9,000 trials per condition", "3.9447 s", "0.3618 s", "10.90x"],
        ["RTX 3060 laptop  9,000 trials per condition", "90.8538 s", "8.6001 s", "10.56x"],
        ["RTX 3060  three settings at 512 trials", "5.06-5.62 s", "0.487-0.576 s", "8.94-10.78x"],
    ], [3.25, 1.2, 1.2, 1.35])
    p(doc, "The October simulator timings cover both Model B conditions. H100 trial chunks were 64; laptop "
      "chunks were 8. The block size was 32. Compilation and warm-up preceded timing, and outputs were "
      "synchronized so the measurements include GPU computation rather than dispatch alone. [2-4]")
    heading(doc, "Trial counts and controls")
    table(doc, ["Check", "Trials and repetitions", "Observed result"], [
        ["H100 paired stopping control", "12 seeds x 2 conditions x 9,000 = 216,000 trial pairs", "Every tested choice and reaction time matched"],
        ["H100 timing", "7 warmed calls per variant at seed 0", "Median speedup 10.90x; 95% timing bootstrap interval 10.83-10.91x"],
        ["Laptop paired stopping control", "3 settings x 3 seeds x 2 conditions x 512 = 9,216 trial pairs", "Every tested choice and reaction time matched"],
        ["Laptop 9,000-trial timing pilot", "3 warmed calls per variant; one seed", "Median speedup 10.56x"],
    ], [1.65, 2.6, 2.75])
    p(doc, "The H100 comparison with the independent original path passed the provisional per-category "
      "choice and RT screens at one parameter setting. The choice margin was one percentage point and the "
      "relative RT-quantile margin was 2%. These are screening margins, not a universal equivalence theorem. "
      "Some laptop category RT checks had too few samples for the statistical screen; exact paired matches "
      "were still verified. The repeat-call timing interval describes timing variability at that configuration, "
      "not uncertainty over every possible parameter setting. [2, 3]")
    heading(doc, "Reproducibility and limits")
    p(doc, "The October runs used JAX 0.10.2 with 64-bit support enabled, 32-bit spatial fields and 64-bit RTs. "
      "The laptop used a two-entry K-L basis cache for both variants to fit its 6 GB memory. "
      "Software, drivers, chunk size and hardware differ between machines; their approximately 23-fold "
      "absolute timing difference is not a pure GPU-silicon comparison. Existing defaults remain unchanged. [2, 4, 5]")

    new_page(doc, "Fitting quality is still unresolved")
    p(doc, "A fast simulator evaluates a candidate parameter vector quickly. A fitter also needs a search "
      "strategy that finds a scientifically acceptable vector. The laptop study compared four paired "
      "starting cases across three true parameter settings, giving eight fits. It used 512 fitting trials "
      "per condition, 9,000 observed trials, and independent scoring at 9,000 trials for three keys. "
      "Each fit allowed up to 1,000 simplex iterations and about 30 training minutes. [4, 5]")
    table(doc, ["Setting and start", "Original minutes", "Early minutes", "Original error", "Early error"], [
        ["Baseline  start 0", "30.08 capped", "9.14", "12.96%", "12.93%"],
        ["Lower boundary  start 0", "29.65", "8.60", "12.71%", "12.52%"],
        ["Higher boundary  start 0", "30.02 capped", "7.88", "12.94%", "11.54%"],
        ["Baseline  start 1", "30.08 capped", "10.35", "9.49%", "7.85%"],
    ], [2.1, 1.35, 1.2, 1.2, 1.15])
    p(doc, "Error is the mean relative error over the twelve active effective parameters; it is not the "
      "percentage of incorrect simulated choices. The unused sv slot remains in the thirteen-coordinate "
      "search but is excluded from recovery error. Training times exclude initialization, warm-up and scoring.", size=9.5)
    picture(doc, ROOT / "docs/results/2026-10-13-rtx3060/fitting-quality.png", 7)
    p(doc, "All eight fits failed both requirements: at most 5% mean active-parameter error and an independent "
      "median G2 no larger than the true-parameter median plus 10% of its absolute magnitude, with a minimum "
      "baseline magnitude of one, plus one unit. Scoring used keys 999, 5001 and 5002 on the same observations "
      "with new simulator randomness. It did not guide fitting and was not held-out participant validation.")
    p(doc, "The early-stop optimizer reported convergence in all four cases, while the original converged "
      "once and was time-capped three times. Nevertheless, neither reached the joint target. "
      "The first start began at 13.05% error; most estimates stayed near it. A different start began at 8.69% "
      "and ended at 7.85% or 9.49%. The June 2.98% example used a different trial budget, dataset, perturbation "
      "and tolerance, so it is not directly contradicted by this study.")

    new_page(doc, "Why laptop fitting stalled")
    p(doc, "A CPU replay reproduced all 3,125 archived optimizer proposals exactly. All five converged fits "
      "ended with zero loss spread across the final simplex, while parameter errors remained above 5%. "
      "The other three fits reached their time caps. A small, flat simplex satisfied the stopping rule; "
      "it did not establish accurate recovery. Final clamping and the recovery-error calculation checked out. [9]")
    p(doc, "The June comparison was easier: its exact starting guess was already at 3.324% mean active-parameter "
      "error and finished at 2.98%. October began at 8.688% or 13.048%. June also used 9,000 fitting trials "
      "per condition; October used 512. June did not establish reliable recovery from October's harder starts.")
    heading(doc, "K-L coefficient assignment")
    p(doc, "The existing basis sorts spatial modes by eigenvalue, while random coefficients are assigned by "
      "column position. A tiny width change can reshuffle modes and attach the same coefficients to different "
      "patterns. On the production CPU basis, width 10 to 10.000001 retained the same 1,325 modes but changed "
      "about 42% of their ranks. The covariance spectrum barely changed, while the same-coefficient field "
      "changed sharply. This mechanism predates improvement 1.")
    picture(doc, ROOT / "docs/results/2026-10-15-model-b-fit-diagnosis/mode-assignment.png", 6.25)
    p(doc, "Aligning coefficients by physical mode removed the demonstrated jump and preserved covariance. "
      "This isolates a problem with comparisons between parameter values, not an invalid Gaussian distribution "
      "at a fixed value. The alignment is a diagnostic intervention, not a deployed or validated fitting repair.", size=10)
    heading(doc, "The small fitting budget")
    p(doc, "Conditional resampling of one saved true-parameter pool gave median G2 scores of 624 at 512 trials "
      "and 85 at 9,000, across 128 replicates per budget. This demonstrates extra score noise, not new fitting "
      "successes. Stable mode-indexed randomness, larger final fitting budgets and restarted searches warrant "
      "controlled tests. No 3060-specific fault was isolated, and no new GPU run was performed. [9]")

    new_page(doc, "What is complete and what needs finalization")
    p(doc, "The implementation and recorded simulation result are complete enough to freeze and describe. "
      "Calling the method fully finalized requires broader scientific validation. We should agree the "
      "parameter domain, tolerances and intended paper claim with Dr Ratcliff before extending the tests.")
    table(doc, ["Evidence area", "Current status", "Completion criterion"], [
        ["Implementation", "Implemented and tested as an optional path", "Freeze a source version and keep passing crossing, padding and no-crossing regression checks"],
        ["Forward simulation fidelity", "Exact paired checks plus one H100 legacy distribution screen", "Cover boundary, noise and drift ranges, including slow, rare and no-crossing cases; use agreed margins and uncertainty"],
        ["Performance", "Large gain on H100 and RTX 3060", "Repeat matched runs on a frozen version; report cold and warmed costs, memory, repetitions and variability"],
        ["Accurate fitting", "Not established; 0 of 8 laptop fits passed", "Use identical observations and starts; test larger fitting budgets and multiple starts; compare time to the same accepted quality"],
        ["Publication contribution", "Engineering result with a proposed methods-paper scope", "Review prior work on early stopping, active-batch compaction and spatial DDM acceleration; define the incremental contribution"],
    ], [1.2, 2.2, 3.6])
    heading(doc, "Recommended next validation sequence")
    p(doc, "First, run a parameter-grid recovery and boundary-case study of the stopping path against its "
      "paired full-horizon control, and statistical comparisons with the original simulator. Use enough "
      "trials in rare response categories to assess the agreed RT and choice margins.")
    p(doc, "A concrete starting proposal is ten agreed parameter settings, twelve seeds and 9,000 trials per "
      "condition: 2.16 million paired trial comparisons across two conditions. Repeat warmed timing seven "
      "times per setting and report first-use costs separately. This proposed study has not been run; "
      "expand it if rare categories or parameter limits remain unresolved.")
    p(doc, "Second, implement and test stable mode-indexed coefficient assignment as a separate experimental "
      "option. Compare it with the original assignment at matched observations, starts and trial budgets, "
      "including a final 9,000-trial optimization stage. Test restarts and removal of the inactive coordinate "
      "separately. Retain failures and time-capped cases; do not average only successful fits. [9]")
    p(doc, "Finally, repeat performance measurements on the frozen implementation and assess real participant "
      "data if it becomes available for this project. A simulation-methods paper can focus on faster forward "
      "simulation with validated output fidelity. A claim of faster accurate fitting needs the additional "
      "recovery result. We can draft the model, algorithm and current benchmark sections now.")
    p(doc, "Other optional extensions reached 0.269 seconds in one separate H100 simulator experiment, "
      "an additional 1.38-fold gain over that run's 0.3714-second early-stop control. Their combined fitting "
      "results were mixed. They are separate exploratory results and do not replace the improvement-1 headline. [8]", size=9.5)

    new_page(doc, "Proposed paper structure")
    for label, text in [
        ("Scientific model", "Define Model B, its spatial noise, decisions, RT calculation and numerical approximations."),
        ("Algorithm and implementation", "Explain blockwise first-crossing detection, stable random indexing, active-trial compaction and GPU execution."),
        ("Validation", "Document paired output checks, independent legacy distribution checks, parameter coverage, margins and numerical precision."),
        ("Performance", "Report matched workload timings, warm-up, memory, hardware, software and uncertainty. Keep historical comparisons separate."),
        ("Fitting and limitations", "Present recovery successes and failures, censored searches, trial budgets and the limits of any time-to-quality claim."),
    ]:
        p(doc, label, bold=True)
        p(doc, text)
    p(doc, "Early stopping, K-L expansions and stream compaction are established ideas. "
      "The potential contribution is their scientifically faithful implementation and evaluation for "
      "this spatial DDM; novelty and publication readiness still require a literature review.")
    heading(doc, "Evidence of record")
    p(doc, "Links 1-8 point to the preserved experiment version; link 9 is the later diagnostic addendum. "
      "These are implementation and benchmark sources, not a substitute for a research bibliography.", size=10)
    for label, path in [
        ("[1] June Stage 6 completion and historical comparison", "docs/notes/2026-06-05-stage-6-completion.md"),
        ("[2] October H100 experiment with timing and fidelity results", "docs/results/2026-10-13-h100/README.md"),
        ("[3] H100 full-budget raw summary and statistical screens", "docs/results/2026-10-13-h100/full-budget/summary.json"),
        ("[4] RTX 3060 complete study report", "docs/results/2026-10-13-rtx3060/README.md"),
        ("[5] Laptop audit of 25 recorded jobs and raw archive hashes", "docs/results/2026-10-13-rtx3060/audit.json"),
        ("[6] Model B simulator and blocked execution entry points", "model_b/simulate.py"),
        ("[7] Existing K-L basis construction and sampling", "model_b/grf_kl.py"),
        ("[8] Separate optional-extension experiment", "docs/results/2026-10-13-h100/EXTENSION_RESULTS.md"),
    ]:
        link(doc, label, path)
    link(doc, "[9] October 15 CPU diagnosis with exact replay and controlled probes",
         "docs/results/2026-10-15-model-b-fit-diagnosis/README.md",
         "https://github.com/anik1617/Ratcliff/blob/feat/h100-blocked-staged-fitting/")
    p(doc, "The H100 and laptop raw archives retain settings, seeds, proposal traces, source snapshots and "
      "SHA-256 manifests. The laptop audit verified all 25 GPU-worker records, all eight returned fit vectors "
      "against their scores, all 9,216 paired trial comparisons, and both archive/file hashes. GitHub archive "
      "contents were independently downloaded and hash-checked after upload. The laptop experiment finished "
      "within the authorized limit; no new GPU experiment was run to prepare these documents.", size=9.5)
    doc.save(OUT / "Model_B_Methods_and_Validation_Report.docx")


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    WORK.mkdir(parents=True, exist_ok=True)
    make_diagrams()
    brief()
    report()
    print(json.dumps({"outputs": [str(p) for p in sorted(OUT.glob("*.docx"))]}, indent=2))
